#!/usr/bin/env python3
"""AssetTiger 내보내기 파일을 이 시스템의 업로드 양식으로 바꾼다.

    python tools/convert_assettiger.py <AssetTiger엑셀> [저장폴더]

결과로 두 개의 파일이 만들어진다.
    1_직원등록.xlsx   - 담당자 목록 (사번은 임시로 생성한다)
    2_자산등록.xlsx   - 자산 목록 (사용자 사번이 채워져 지급 처리까지 된다)

직원 파일을 먼저 올린 뒤 자산 파일을 올려야 지급 이력이 함께 만들어진다.

AssetTiger 는 사번을 내보내지 않으므로 사번을 임시로 만든다.
실제 사번이 있다면 1_직원등록.xlsx 의 '사번' 열을 바꾸고, 2_자산등록.xlsx 의
'사용자 사번' 열도 같이 바꾼 뒤 올리면 된다.
"""

from __future__ import annotations

import sys
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

# --- AssetTiger 값 → 이 시스템의 값 ------------------------------------------

STATUS_MAP = {
    "checked out": "사용중",
    "available": "재고",
    "under repair": "수리중",
    "lost/missing": "분실",
    "lost": "분실",
    "missing": "분실",
    "disposed": "폐기",
    "broken": "수리중",
    "sold": "폐기",
    "donated": "폐기",
}

# 자산 설명에서 분류를 추정한다. 먼저 걸리는 것을 쓴다.
CATEGORY_RULES = [
    (("노트북", "랩탑", "laptop", "notebook"), "노트북"),
    (("데스크탑", "데스크톱", "desktop", "pc", "본체"), "데스크톱"),
    (("모니터", "monitor", "디스플레이"), "모니터"),
    (("서버", "server", "nas"), "서버"),
    (("스위치", "공유기", "라우터", "방화벽", "switch", "router", "ap"), "네트워크 장비"),
    (("폰", "phone", "태블릿", "tablet", "ipad", "아이패드"), "모바일 기기"),
    (("프린터", "printer", "복합기", "스캐너", "키보드", "마우스", "도킹", "프로젝터"), "주변기기"),
    (("라이선스", "license", "office", "windows", "소프트웨어", "software"), "소프트웨어"),
]

# 담당자 칸에 사람이 아니라 장소가 들어가 있는 경우를 가려낸다.
PLACE_HINTS = ("사무실", "센터", "창고", "지점", "본사", "office", "실", "룸", "room")


def guess_category(description: str) -> str:
    text = (description or "").lower()
    for keywords, label in CATEGORY_RULES:
        if any(word in text for word in keywords):
            return label
    return "기타"


def looks_like_place(name: str) -> bool:
    """담당자 칸의 값이 사람 이름이 아니라 장소로 보이는지."""
    if not name:
        return False
    return any(hint in name.lower() for hint in PLACE_HINTS)


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def parse_money(value) -> float | None:
    text = clean(value).replace(",", "").replace("₩", "").replace("원", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def read_rows(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    headers = [clean(c.value) for c in ws[1]]
    rows = []
    for raw in ws.iter_rows(min_row=2):
        values = [c.value for c in raw]
        if not any(v is not None and clean(v) for v in values):
            continue
        rows.append(dict(zip(headers, values)))
    wb.close()
    return rows


ASSET_HEADERS = [
    "자산번호", "자산명", "분류", "상태", "제조사", "모델명", "시리얼번호", "사양",
    "보관위치", "공급업체", "도입일", "취득가액", "보증만료일", "라이선스키",
    "사용자 사번", "사용자명", "상각방법", "내용연수", "잔존가치", "장부가액", "비고",
]
EMPLOYEE_HEADERS = ["사번", "이름", "부서", "직급", "이메일", "연락처", "재직상태", "비고"]

HEADER_FILL = "1F3B63"


def _sheet(wb, title, headers, widths):
    ws = wb.active if wb.active.max_row == 1 and wb.active.max_column == 1 else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    return ws


def convert(source: Path, out_dir: Path) -> dict:
    rows = read_rows(source)
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1) 담당자 모으기 (이름 + 부서). 장소로 보이는 값은 직원으로 만들지 않는다.
    people: OrderedDict[str, str] = OrderedDict()   # 이름 -> 부서
    places: list[str] = []
    for row in rows:
        name = clean(row.get("Assigned to"))
        if not name:
            continue
        if looks_like_place(name):
            if name not in places:
                places.append(name)
            continue
        dept = clean(row.get("Department"))
        # 같은 사람이 여러 행에 있고 부서가 빈 행이 섞여 있으면, 채워진 값을 쓴다
        if name not in people or (dept and not people[name]):
            people[name] = dept

    emp_no_of = {name: f"E{index:04d}" for index, name in enumerate(sorted(people), start=1)}

    # 2) 직원 파일
    wb = Workbook()
    ws = _sheet(wb, "직원등록", EMPLOYEE_HEADERS, [14, 12, 16, 12, 26, 16, 10, 30])
    for name in sorted(people):
        ws.append([emp_no_of[name], name, people[name], "", "", "", "재직", "AssetTiger 에서 이관"])
    employee_path = out_dir / "1_직원등록.xlsx"
    wb.save(employee_path)

    # 3) 자산 파일
    wb = Workbook()
    ws = _sheet(wb, "자산등록", ASSET_HEADERS,
                [14, 20, 14, 10, 14, 20, 20, 24, 18, 14, 12, 14, 12, 20, 14, 12, 12, 10, 12, 14, 36])

    place_assigned: list[str] = []
    unknown_status: set[str] = set()

    for row in rows:
        tag = clean(row.get("Asset Tag ID"))
        description = clean(row.get("Description")) or "자산"
        raw_status = clean(row.get("Status"))
        status = STATUS_MAP.get(raw_status.lower(), "재고")
        if raw_status and raw_status.lower() not in STATUS_MAP:
            unknown_status.add(raw_status)

        assignee = clean(row.get("Assigned to"))
        notes = clean(row.get("Event Notes"))

        location = ""
        emp_no = ""
        holder_name = ""

        if assignee and looks_like_place(assignee):
            # 사람이 아니라 장소에 배정된 자산: 보관위치로 옮기고 미지급으로 둔다
            location = assignee
            status = "재고"
            place_assigned.append(tag)
        elif assignee:
            emp_no = emp_no_of.get(assignee, "")
            holder_name = assignee

        # 지급 대상이 없으면 '사용중'일 수 없다
        if not emp_no and status == "사용중":
            status = "재고"

        note_parts = [part for part in (notes, f"AssetTiger 이관 (원래 상태: {raw_status})") if part]

        ws.append([
            tag,
            description,
            guess_category(description),
            status,
            "",                                   # 제조사 - AssetTiger 내보내기에 없음
            clean(row.get("Model")),
            clean(row.get("Serial No")),
            "",                                   # 사양
            location,
            "",                                   # 공급업체
            clean(row.get("Purchase Date")),
            parse_money(row.get("Cost")),
            "",                                   # 보증만료일
            "",                                   # 라이선스키
            emp_no,
            holder_name,
            "", "", "", "",                       # 감가상각 4개 열은 비워 둔다
            " / ".join(note_parts),
        ])

    for line in ws.iter_rows(min_row=2):
        line[10].number_format = "yyyy-mm-dd"
        line[11].number_format = "#,##0"
    asset_path = out_dir / "2_자산등록.xlsx"
    wb.save(asset_path)

    return {
        "rows": len(rows),
        "employees": len(people),
        "places": places,
        "place_assigned": place_assigned,
        "unknown_status": sorted(unknown_status),
        "employee_path": employee_path,
        "asset_path": asset_path,
    }


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2

    source = Path(argv[1])
    out_dir = Path(argv[2]) if len(argv) > 2 else source.parent / "변환결과"
    if not source.exists():
        print(f"파일을 찾을 수 없습니다: {source}")
        return 1

    info = convert(source, out_dir)

    print(f"\n  AssetTiger 자산 {info['rows']}건을 변환했습니다.\n")
    print(f"    1) {info['employee_path']}   직원 {info['employees']}명")
    print(f"    2) {info['asset_path']}   자산 {info['rows']}건")
    print("\n  이 순서대로 올려야 지급 이력이 함께 만들어집니다.")

    if info["places"]:
        print(f"\n  [확인 필요] 담당자 칸에 장소로 보이는 값이 있었습니다: {', '.join(info['places'])}")
        print(f"             해당 자산({len(info['place_assigned'])}건)은 보관위치로 옮기고 '재고'로 뒀습니다.")
        print(f"             자산번호: {', '.join(info['place_assigned'])}")
    if info["unknown_status"]:
        print(f"\n  [확인 필요] 모르는 상태값이 있어 '재고'로 처리했습니다: {', '.join(info['unknown_status'])}")

    print("\n  사번은 임시로 만든 값(E0001...)입니다.")
    print("  실제 사번이 있다면 두 파일의 '사번' / '사용자 사번' 열을 함께 바꾼 뒤 올리세요.\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
