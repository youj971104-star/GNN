"""직원 엑셀 업로드 테스트."""

import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.excel import EMPLOYEE_COLUMNS, export_employee_template, import_employees
from app.models import Employee

HEADERS = [name for name, _ in EMPLOYEE_COLUMNS]


def _file(rows: list[list], headers: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers or HEADERS)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _row(emp_no="E001", name="홍길동", **cells) -> list:
    values = {"사번": emp_no, "이름": name, "부서": "개발팀", "재직상태": "재직"}
    values.update(cells)
    return [values.get(header) for header in HEADERS]


def test_직원을_일괄_등록한다(db):
    result = import_employees(db, _file([_row(), _row("E002", "김철수", 부서="영업팀")]))

    assert result.created == 2 and result.skipped == 0
    emp = db.scalar(select(Employee).where(Employee.emp_no == "E001"))
    assert emp.name == "홍길동"
    assert emp.department == "개발팀"
    assert emp.status == "ACTIVE"


def test_같은_사번은_갱신한다(db, employee):
    result = import_employees(db, _file([_row(employee.emp_no, "이름 변경됨", 부서="총무팀")]))

    assert result.created == 0 and result.updated == 1
    db.expire_all()
    assert db.get(Employee, employee.id).name == "이름 변경됨"
    assert db.get(Employee, employee.id).department == "총무팀"


def test_사번이나_이름이_없으면_건너뛴다(db):
    result = import_employees(db, _file([
        _row("E001", "정상"),
        _row("", "사번없음"),
        _row("E003", ""),
    ]))
    assert result.created == 1
    assert result.skipped == 2
    assert len(result.errors) == 2


def test_한_파일에_같은_사번이_두_번_나오면_오류(db):
    result = import_employees(db, _file([_row("E001", "첫번째"), _row("E001", "두번째")]))
    assert result.created == 1 and result.skipped == 1
    assert "중복" in result.errors[0]


def test_이메일_형식을_검사한다(db):
    result = import_employees(db, _file([_row(이메일="잘못된주소")]))
    assert result.skipped == 1
    assert "이메일" in result.errors[0]


def test_재직상태를_한글로_적을_수_있다(db):
    import_employees(db, _file([
        _row("E001", "재직자", 재직상태="재직"),
        _row("E002", "퇴사자", 재직상태="퇴사"),
        _row("E003", "기본값", 재직상태=None),
    ]))
    statuses = {e.emp_no: e.status for e in db.scalars(select(Employee)).all()}
    assert statuses["E001"] == "ACTIVE"
    assert statuses["E002"] == "RESIGNED"
    assert statuses["E003"] == "ACTIVE"  # 비우면 재직


def test_필수_열이_없으면_전체를_거부한다(db):
    result = import_employees(db, _file([["값"]], headers=["엉뚱한열"]))
    assert result.created == 0
    assert "필수 열이 없습니다" in result.errors[0]


def test_엑셀이_아니면_안내한다(db):
    result = import_employees(db, b"not an excel file")
    assert result.has_errors
    assert "엑셀 파일을 열 수 없습니다" in result.errors[0]


def test_양식을_내려받아_그대로_채워_올릴_수_있다(db):
    wb = load_workbook(io.BytesIO(export_employee_template()))
    ws = wb.active
    assert [c.value for c in ws[1]] == HEADERS

    # 예시 행을 실제 값으로 바꿔 올린다
    ws.delete_rows(2)
    ws.append(["2026001", "신입사원", "개발팀", "사원", "new@example.com", "010-0000-0000", "재직", ""])
    buffer = io.BytesIO()
    wb.save(buffer)

    result = import_employees(db, buffer.getvalue())
    assert result.created == 1 and result.skipped == 0


def test_직원을_올린_뒤_자산에서_지급_처리까지_이어진다(db):
    """이관 시나리오: 직원 먼저, 그다음 자산."""
    from app.excel import ASSET_COLUMNS, import_assets
    from app.models import Asset, Assignment

    import_employees(db, _file([_row("E100", "김지급", 부서="개발팀")]))

    asset_headers = [label for _, label, _ in ASSET_COLUMNS]
    values = {"자산번호": "IT-1", "자산명": "노트북", "분류": "노트북",
              "상태": "사용중", "사용자 사번": "E100"}
    wb = Workbook()
    ws = wb.active
    ws.append(asset_headers)
    ws.append([values.get(h) for h in asset_headers])
    buffer = io.BytesIO()
    wb.save(buffer)

    result = import_assets(db, buffer.getvalue(), actor="admin")
    assert result.created == 1 and result.skipped == 0

    asset = db.scalar(select(Asset).where(Asset.asset_no == "IT-1"))
    assert asset.status == "IN_USE"
    assert asset.holder.name == "김지급"
    assert db.scalar(select(Assignment).where(Assignment.asset_id == asset.id)) is not None


def test_화면에서_직원_양식과_업로드가_동작한다(admin_client):
    template = admin_client.get("/employees/template")
    assert template.status_code == 200
    assert template.content[:2] == b"PK"

    page = admin_client.get("/employees/import")
    assert page.status_code == 200

    rejected = admin_client.post(
        "/employees/import", files={"file": ("a.csv", b"a,b", "text/csv")}
    )
    assert rejected.status_code == 400


def test_일반_사용자는_직원_업로드에_접근할_수_없다(viewer_client):
    assert viewer_client.get("/employees/import").status_code == 403


def test_사번을_수정할_수_있다(admin_client, db, employee):
    """이관 시 임시 사번이 들어갈 수 있어, 나중에 실제 사번으로 바꿀 수 있어야 한다."""
    response = admin_client.post(
        f"/employees/{employee.id}/edit",
        data={"emp_no": "2026999", "name": employee.name, "status": "ACTIVE"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    db.expire_all()
    assert db.get(Employee, employee.id).emp_no == "2026999"


def test_다른_직원이_쓰는_사번으로는_바꿀_수_없다(admin_client, db, employee):
    db.add(Employee(emp_no="E999", name="다른사람"))
    db.commit()

    response = admin_client.post(
        f"/employees/{employee.id}/edit",
        data={"emp_no": "E999", "name": employee.name, "status": "ACTIVE"},
    )
    assert response.status_code == 400
    assert "이미 다른 직원" in response.text
