"""목록 칸 제목을 눌러 정렬하기."""

import io
import re
from datetime import date

import pytest
from openpyxl import load_workbook

from app.models import ASSET_CATEGORIES, Asset, Employee, Maintenance
from app.services import assign_asset, return_asset
from app.sorting import SortColumn, read_sort

# --- 주소 읽기 -------------------------------------------------------------------

COLUMNS = {
    "name": SortColumn(("name",)),
    "price": SortColumn(("price",), default_desc=True),
}


def test_아무것도_없으면_기본_정렬():
    sort = read_sort({}, COLUMNS, "name")
    assert (sort.key, sort.direction) == ("name", "asc")


@pytest.mark.parametrize("params", [{"sort": "password_hash"}, {"sort": "__class__"}, {"sort": ""}])
def test_모르는_칸은_기본_정렬로_되돌린다(params):
    assert read_sort(params, COLUMNS, "name").key == "name"


def test_방향이_이상하면_그_칸의_기본_방향():
    assert read_sort({"sort": "price", "dir": "sideways"}, COLUMNS, "name").direction == "desc"
    assert read_sort({"sort": "name", "dir": "DROP"}, COLUMNS, "name").direction == "asc"


def test_누르면_오름차순부터_같은_칸을_다시_누르면_뒤집힌다():
    sort = read_sort({"sort": "name", "dir": "asc"}, COLUMNS, "name")
    assert sort.next_direction("name") == "desc"
    assert sort.next_direction("price") == "asc"        # 다른 칸은 오름차순부터

    sort = read_sort({"sort": "name", "dir": "desc"}, COLUMNS, "name")
    assert sort.next_direction("name") == "asc"
    assert sort.state("name") == "descending"
    assert sort.state("price") == "none"


# --- 자산 목록 -------------------------------------------------------------------

def _rows(html: str) -> list[str]:
    """목록 표의 자산번호를 위에서부터."""
    return re.findall(r'data-href="/assets/\d+">\s*<td class="mono nowrap"><a [^>]*>([^<]+)</a>', html)


@pytest.fixture
def assets(db):
    holders = [Employee(emp_no="E1", name="나영희"), Employee(emp_no="E2", name="가철수")]
    items = [
        Asset(asset_no="A-1", name="모니터", category="MONITOR", purchase_price=300_000,
              purchase_date=date(2024, 1, 1), location="본사"),
        Asset(asset_no="A-2", name="노트북", category="NOTEBOOK", purchase_price=2_000_000,
              purchase_date=date(2025, 1, 1), location="창고"),
        Asset(asset_no="A-3", name="데스크톱", category="DESKTOP", purchase_price=None,
              purchase_date=None, location=None),
        Asset(asset_no="A-4", name="서버", category="SERVER", purchase_price=9_000_000,
              purchase_date=date(2023, 1, 1), location="전산실"),
    ]
    db.add_all(holders + items)
    db.commit()
    assign_asset(db, asset=items[0], employee=holders[0])   # A-1 → 나영희
    assign_asset(db, asset=items[3], employee=holders[1])   # A-4 → 가철수
    db.commit()
    return items


@pytest.mark.parametrize(
    "key,direction,expected",
    [
        ("asset_no", "asc", ["A-1", "A-2", "A-3", "A-4"]),
        ("asset_no", "desc", ["A-4", "A-3", "A-2", "A-1"]),
        # 값이 없는 행(A-3)은 방향과 상관없이 맨 뒤
        ("purchase_price", "asc", ["A-1", "A-2", "A-4", "A-3"]),
        ("purchase_price", "desc", ["A-4", "A-2", "A-1", "A-3"]),
        ("purchase_date", "asc", ["A-4", "A-1", "A-2", "A-3"]),
        ("location", "asc", ["A-1", "A-4", "A-2", "A-3"]),          # 본사 < 전산실 < 창고
        # 사용자 이름 순, 미지급은 맨 뒤
        ("holder", "asc", ["A-4", "A-1", "A-2", "A-3"]),           # 가철수 < 나영희
        ("holder", "desc", ["A-1", "A-4", "A-2", "A-3"]),
    ],
)
def test_자산_목록을_칸마다_정렬한다(admin_client, assets, key, direction, expected):
    html = admin_client.get(f"/assets?sort={key}&dir={direction}").text
    assert _rows(html) == expected


def test_분류는_검색_목록에_나오는_순서를_따른다(admin_client, assets):
    """코드의 알파벳 순서(DESKTOP, MONITOR …)가 아니라 노트북·데스크톱·모니터… 순서."""
    html = admin_client.get("/assets?sort=category&dir=asc").text
    order = list(ASSET_CATEGORIES)
    got = [a.category for a in sorted(assets, key=lambda a: order.index(a.category))]
    by_no = {a.asset_no: a.category for a in assets}
    assert [by_no[no] for no in _rows(html)] == got


def test_정렬_중인_칸_제목에_표시가_붙고_누르면_뒤집힌다(admin_client, assets):
    html = admin_client.get("/assets?sort=purchase_price&dir=asc").text
    assert 'aria-sort="ascending"' in html
    assert "sort=purchase_price&amp;dir=desc" in html     # 다시 누르면 내림차순
    assert "sort=asset_no&amp;dir=asc" in html            # 다른 칸은 오름차순부터


def test_검색해도_정렬이_유지된다(admin_client, assets):
    html = admin_client.get("/assets?sort=purchase_price&dir=desc").text
    assert '<input type="hidden" name="sort" value="purchase_price">' in html
    assert '<input type="hidden" name="dir" value="desc">' in html


def test_정렬을_바꾸면_첫_쪽으로_돌아간다(admin_client, assets):
    html = admin_client.get("/assets?page=3&sort=asset_no&dir=asc").text
    links = re.findall(r'href="(\?[^"]*sort=[^"]*)"', html)
    assert links and all("page=" not in link for link in links)


def test_예전_주소도_그대로_동작한다(admin_client, assets):
    """방향이 없던 시절의 ?sort=purchase_price 는 '큰 값이 위' 였다."""
    html = admin_client.get("/assets?sort=purchase_price").text
    assert _rows(html)[0] == "A-4"


def test_엑셀로_내려받아도_같은_순서다(admin_client, assets):
    response = admin_client.get("/assets/export?sort=purchase_price&dir=desc")
    sheet = load_workbook(io.BytesIO(response.content)).active
    numbers = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    assert numbers == ["A-4", "A-2", "A-1", "A-3"]


# --- 다른 목록 -------------------------------------------------------------------

def test_직원은_보유_자산_수로_정렬한다(admin_client, assets):
    html = admin_client.get("/employees?sort=assets&dir=desc").text
    assert 'aria-sort="descending"' in html
    names = re.findall(r'data-href="/employees/\d+">.*?<td>([^<]+)</td>', html, re.S)
    assert names[:2] in (["나영희", "가철수"], ["가철수", "나영희"])     # 둘 다 1대


def test_직원_이름_정렬(admin_client, assets):
    html = admin_client.get("/employees?sort=name&dir=desc").text
    first = html.index("나영희")
    assert first < html.index("가철수")


def test_지급_이력은_보유일수로_정렬한다(admin_client, db, assets):
    holder = db.query(Employee).filter_by(emp_no="E1").one()
    old = Asset(asset_no="A-9", name="오래된 노트북")
    db.add(old); db.commit()
    assign_asset(db, asset=old, employee=holder, assigned_at=date(2020, 1, 1))
    db.commit()
    return_asset(db, assignment=old.assignments[0], returned_at=date(2020, 1, 11))
    db.commit()

    # A-9 는 10일 쓰고 반납, A-1 은 오늘 지급해 0일째
    longest_first = admin_client.get("/assignments?sort=days&dir=desc").text
    shortest_first = admin_client.get("/assignments?sort=days&dir=asc").text
    assert longest_first.index("A-9") < longest_first.index("A-1")
    assert shortest_first.index("A-1") < shortest_first.index("A-9")


def test_정비_이력은_비용으로_정렬한다(admin_client, db, assets):
    db.add_all([
        Maintenance(asset_id=assets[0].id, maintained_at=date(2026, 1, 1), kind="REPAIR",
                    description="싼 수리", cost=10_000),
        Maintenance(asset_id=assets[1].id, maintained_at=date(2026, 2, 1), kind="INSPECTION",
                    description="비싼 수리", cost=500_000),
    ])
    db.commit()
    html = admin_client.get("/maintenance?sort=cost&dir=desc").text
    assert html.index("비싼 수리") < html.index("싼 수리")
    html = admin_client.get("/maintenance?sort=cost&dir=asc").text
    assert html.index("싼 수리") < html.index("비싼 수리")


@pytest.mark.parametrize(
    "path,keys",
    [
        ("/assets", ["asset_no", "name", "category", "status", "holder", "location",
                     "purchase_date", "purchase_price"]),
        ("/employees", ["emp_no", "name", "department", "position", "email", "phone",
                        "status", "assets"]),
        ("/assignments", ["asset_no", "asset_name", "employee", "department", "assigned_at",
                          "returned_at", "days", "state"]),
        ("/maintenance", ["maintained_at", "asset_no", "asset_name", "kind", "vendor",
                          "description", "cost", "next_due"]),
    ],
)
def test_모든_칸을_양쪽_방향으로_정렬할_수_있다(admin_client, assets, path, keys):
    for key in keys:
        for direction, state in (("asc", "ascending"), ("desc", "descending")):
            response = admin_client.get(f"{path}?sort={key}&dir={direction}")
            assert response.status_code == 200, (path, key, direction)
            assert f'aria-sort="{state}"' in response.text, (path, key, direction)
