"""인력현황 엑셀을 직원 명부에 반영하기.

실제 인력현황과 같은 모양(묶음 제목 '* 협력사' · '* 퇴사자', 묶음마다 다시 나오는 머리글,
'-' 로 비운 칸, 시트마다 한 시점)을 가짜 이름으로 만들어 시험한다.
"""

import io
import re
from datetime import datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app import roster
from app.models import Asset, Employee

HEADER = [None, "번호", "구분", "부서", "팀", "직급(책)", "성명", "고용형태", "입사일", "비고"]


def _sheet(ws, main, partners=(), resigned=()):
    ws.append([])
    ws.append([])
    ws.append([None, "■ ㈜테스트 인력현황_파견 포함", None, None, None, None, None, None, datetime(2026, 9, 1)])
    ws.append(HEADER)
    for i, (corp, dept, name, *rest) in enumerate(main, start=1):
        note = rest[0] if rest else None
        ws.append([None, i, corp, dept, dept, "매니저", name, "정규", datetime(2024, 1, 1), note])
    if partners:
        ws.append([])
        ws.append([None, "* 협력사"])
        ws.append(HEADER)
        for i, (corp, dept, name) in enumerate(partners, start=1):
            ws.append([None, i, corp, dept, None, "-", name, "협력업체", None, None])
    if resigned:
        ws.append([])
        ws.append([None, "* 퇴사자"])
        ws.append([None, "번호", "구분", "부서", "팀", "직급(책)", "성명", "고용형태", "퇴사일", "비고"])
        for i, (corp, dept, name) in enumerate(resigned, start=1):
            ws.append([None, i, corp, dept, dept, "매니저", name, "계약", datetime(2026, 10, 1), None])


def _file(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, kwargs in sheets.items():
        _sheet(wb.create_sheet(title), **kwargs)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


LATEST = dict(
    main=[
        ("가법인", "CS1팀", "김하나"),
        ("나법인", "QA1팀", "이둘"),
        ("가법인", "CS2팀", "박셋", "육아휴직(26.01.01-26.12.31)"),
        ("나법인", "HR팀", "최넷"),
        ("나법인", "HR팀", "정다섯"),
    ],
    partners=[("위탁계약", "제주", "한여섯")],
    resigned=[("가법인", "QA2팀", "오일곱"), ("나법인", "QA2팀", "명부에없음")],
)


# --- 읽기 ------------------------------------------------------------------------

def test_묶음과_머리글을_읽는다():
    parsed = roster.read_roster(_file({"1001": dict(main=[("가법인", "A", "옛사람")]), "1019": LATEST}))
    assert parsed.sheet == "1019"                       # 맨 오른쪽 = 가장 최근
    assert parsed.sheets == ["1001", "1019"]
    assert [p.name for p in parsed.active] == ["김하나", "이둘", "박셋", "최넷", "정다섯", "한여섯"]
    assert [p.name for p in parsed.resigned] == ["오일곱", "명부에없음"]
    assert dict(parsed.group_counts()) == {"본 인력현황": 5, "협력사": 1, "퇴사자": 2}


def test_시트를_고를_수_있다():
    content = _file({"1001": dict(main=[("가법인", "A", "옛사람")]), "1019": LATEST})
    assert [p.name for p in roster.read_roster(content, "1001").people] == ["옛사람"]


def test_칸_값을_정리한다():
    parsed = roster.read_roster(_file({"s": LATEST}))
    partner = next(p for p in parsed.people if p.name == "한여섯")
    assert partner.company == "위탁계약" and partner.group == "협력사"
    assert partner.position is None                     # '-' 는 빈칸
    kim = parsed.people[0]
    assert kim.joined == "2024-01-01" and kim.kind == "정규"
    assert next(p for p in parsed.people if p.name == "박셋").on_leave


def test_인력현황이_아니면_알려_준다():
    wb = Workbook()
    wb.active.append(["사번", "이름"])
    buffer = io.BytesIO()
    wb.save(buffer)
    with pytest.raises(roster.RosterError, match="인력현황 시트를 찾지 못했습니다"):
        roster.read_roster(buffer.getvalue())
    with pytest.raises(roster.RosterError, match="엑셀 파일을 열 수 없습니다"):
        roster.read_roster(b"not excel")


@pytest.mark.parametrize("name", ["홍길동", "홍 길동", " 홍길동 ", "홍길동(QA1팀)", "홍길동 [파견]", "홍길동（퇴사예정）"])
def test_이름은_공백과_괄호_설명을_빼고_비교한다(name):
    assert roster.name_key(name) == "홍길동"


# --- 미리보기와 반영 -------------------------------------------------------------

@pytest.fixture
def staff(db):
    people = {
        "김하나": Employee(emp_no="E0001", name="김하나", department="CS1팀"),
        "이 둘": Employee(emp_no="E0002", name="이 둘", department="옛부서", position="선임"),  # 띄어쓰기
        "최넷": Employee(emp_no="E0003", name="최넷", company="나법인"),                      # 이미 맞음
        "정다섯": Employee(emp_no="E0004", name="정다섯", status="RESIGNED"),                 # 다시 재직
        "오일곱": Employee(emp_no="E0005", name="오일곱"),                                   # 퇴사자 명단
        "떠난사람": Employee(emp_no="E0006", name="떠난사람"),                               # 인력현황에 없음
        "이미퇴사": Employee(emp_no="E0007", name="이미퇴사", status="RESIGNED"),
    }
    db.add_all(people.values())
    db.commit()
    db.add(Asset(asset_no="A-1", name="노트북", status="IN_USE", holder_id=people["떠난사람"].id))
    db.commit()
    return people


def test_미리보기(db, staff):
    plan = roster.plan_sync(db, roster.read_roster(_file({"s": LATEST})))
    assert sorted(m.employee.name for m in plan.company_changes) == ["김하나", "이 둘"]
    # 법인은 이미 같지만 팀·고용형태·입사일이 비어 있어 채워진다
    info = {m.employee.name: m.changed for m in plan.info_changes}
    assert info == {"최넷": ["팀", "고용형태", "입사일"]}
    assert plan.unchanged == []
    assert [m.employee.name for m in plan.reactivate] == ["정다섯"]
    assert sorted(p.name for p in plan.add) == ["박셋", "한여섯"]
    resign = {r.employee.name: (r.reason, r.holding) for r in plan.resign}
    assert resign == {"떠난사람": ("인력현황에 없음", 1), "오일곱": ("인력현황 퇴사자 명단", 0)}
    assert plan.resign[0].employee.name == "떠난사람"     # 자산을 가진 사람이 먼저
    assert [p.name for p in plan.unknown_resigned] == ["명부에없음"]
    assert plan.ambiguous == []


def test_반영(db, staff):
    result = roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})))
    assert (result.company, result.added, result.resigned, result.reactivated) == (2, 2, 2, 1)
    assert result.resigned_holding == 1

    def emp(name):
        return db.scalar(select(Employee).where(Employee.name == name))

    lee = emp("이 둘")
    # 겹치는 사람은 법인만 바뀐다
    assert (lee.company, lee.department, lee.position) == ("나법인", "옛부서", "선임")
    assert emp("떠난사람").status == "RESIGNED"
    assert emp("오일곱").status == "RESIGNED" and emp("오일곱").company == "가법인"
    assert emp("정다섯").status == "ACTIVE" and emp("정다섯").company == "나법인"
    assert emp("이미퇴사").status == "RESIGNED"

    park = emp("박셋")
    assert park.status == "LEAVE" and park.company == "가법인" and park.department == "CS2팀"
    assert park.note == "육아휴직(26.01.01-26.12.31)"            # 인력현황의 비고 그대로
    assert (park.team, park.employment_type, str(park.hired_on)) == ("CS2팀", "정규", "2024-01-01")
    assert park.roster_group is None and park.roster_order is not None
    partner = emp("한여섯")
    assert partner.company == "위탁계약" and partner.roster_group == "협력사"
    assert partner.employment_type == "협력업체" and partner.position is None
    # 새 사번은 기존 번호 뒤로 이어진다
    assert sorted([park.emp_no, partner.emp_no]) == ["E0008", "E0009"]


def test_두_번_반영해도_더_바뀌지_않는다(db, staff):
    content = _file({"s": LATEST})
    roster.apply_sync(db, roster.read_roster(content))
    plan = roster.plan_sync(db, roster.read_roster(content))
    assert not plan.has_changes


def test_동명이인은_건드리지_않는다(db):
    db.add_all([Employee(emp_no="E1", name="김하나", department="A"),
                Employee(emp_no="E2", name="김하나", department="B")])
    db.commit()
    content = _file({"s": dict(main=[("가법인", "CS1팀", "김하나")])})
    plan = roster.plan_sync(db, roster.read_roster(content))
    assert [a.name for a in plan.ambiguous] == ["김하나"]
    assert plan.resign == [] and plan.add == [] and plan.company_changes == []
    roster.apply_sync(db, roster.read_roster(content))
    assert {e.status for e in db.scalars(select(Employee)).all()} == {"ACTIVE"}


def test_인력현황_안의_동명이인도_건드리지_않는다(db):
    db.add(Employee(emp_no="E1", name="김하나"))
    db.commit()
    content = _file({"s": dict(main=[("가법인", "A", "김하나"), ("나법인", "B", "김하나")])})
    plan = roster.plan_sync(db, roster.read_roster(content))
    assert len(plan.ambiguous) == 1 and not plan.has_changes


# --- 화면 ------------------------------------------------------------------------

def _preview(client, content: bytes):
    return client.post(
        "/employees/roster",
        files={"file": ("인력현황.xlsx", content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_화면에서_미리보고_반영한다(admin_client, db, staff):
    page = _preview(admin_client, _file({"s": LATEST}))
    assert page.status_code == 200
    assert "법인 입력" in page.text and "떠난사람" in page.text and "회수 필요" in page.text
    token = re.search(r'name="token" value="([^"]+)"', page.text).group(1)

    done = admin_client.post("/employees/roster/apply", data={"token": token}, follow_redirects=False)
    assert done.status_code == 303
    db.expire_all()
    assert db.scalar(select(Employee).where(Employee.name == "박셋")) is not None
    assert "인력현황을 반영했습니다" in admin_client.get("/employees").text


def test_명단을_고쳐_보내면_받지_않는다(admin_client, staff):
    page = _preview(admin_client, _file({"s": LATEST}))
    token = re.search(r'name="token" value="([^"]+)"', page.text).group(1)
    forged = token[:-3] + ("AAA" if not token.endswith("AAA") else "BBB")
    assert admin_client.post("/employees/roster/apply", data={"token": forged}).status_code == 400


def test_엑셀이_아니면_안내(admin_client):
    response = admin_client.post("/employees/roster", files={"file": ("a.txt", b"x", "text/plain")})
    assert response.status_code == 400 and ".xlsx" in response.text


def test_관리자만_쓸_수_있다(viewer_client):
    assert viewer_client.get("/employees/roster", follow_redirects=False).status_code == 403
    assert viewer_client.post("/employees/roster/apply", data={"token": "x"},
                              follow_redirects=False).status_code == 403


# --- 법인이 보이는 곳 ------------------------------------------------------------

def test_직원_목록에_법인_칸과_필터(admin_client, db):
    db.add_all([Employee(emp_no="E1", name="가", company="가법인"),
                Employee(emp_no="E2", name="나", company="나법인")])
    db.commit()
    html = admin_client.get("/employees").text
    assert '<option value="가법인"' in html and "sort=company" in html
    filtered = admin_client.get("/employees?company=나법인").text
    assert ">나<" in filtered and ">가<" not in filtered


def test_직원_등록_화면에서_법인을_고른다(admin_client, db):
    db.add(Employee(emp_no="E1", name="가", company="가법인"))
    db.commit()
    assert '<option value="가법인">' in admin_client.get("/employees/new").text
    admin_client.post("/employees/new", data={"emp_no": "E2", "name": "새사람", "company": "나법인",
                                              "status": "ACTIVE"})
    assert db.scalar(select(Employee).where(Employee.emp_no == "E2")).company == "나법인"


def test_자산을_사용자의_법인으로_거른다(admin_client, db):
    a, b = Employee(emp_no="E1", name="가", company="가법인"), Employee(emp_no="E2", name="나", company="나법인")
    db.add_all([a, b])
    db.commit()
    db.add_all([Asset(asset_no="X-1", name="가의 노트북", status="IN_USE", holder_id=a.id),
                Asset(asset_no="X-2", name="나의 노트북", status="IN_USE", holder_id=b.id)])
    db.commit()
    html = admin_client.get("/assets?company=가법인").text
    assert "가의 노트북" in html and "나의 노트북" not in html and "법인: 가법인" in html


def test_대시보드에_법인별_보유_현황(admin_client, db):
    a = Employee(emp_no="E1", name="가", company="가법인")
    db.add(a)
    db.commit()
    db.add(Asset(asset_no="X-1", name="노트북", status="IN_USE", holder_id=a.id))
    db.commit()
    html = admin_client.get("/").text
    assert "법인별 보유 현황" in html and "/assets?company=" in html


# --- 인력현황 항목과 퇴사일 ------------------------------------------------------

def test_기존_직원도_팀_고용형태_입사일을_채우고_부서_직급은_그대로(db, staff):
    staff["김하나"].note = "기존 메모"
    db.commit()
    roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})))
    kim = db.scalar(select(Employee).where(Employee.name == "김하나"))
    assert (kim.team, kim.employment_type, str(kim.hired_on)) == ("CS1팀", "정규", "2024-01-01")
    assert kim.department == "CS1팀" and kim.note == "기존 메모"     # 비고는 비었을 때만 채운다
    lee = db.scalar(select(Employee).where(Employee.name == "이 둘"))
    assert (lee.department, lee.position) == ("옛부서", "선임")


def test_퇴사일을_남긴다(db, staff):
    from datetime import date
    staff["정다섯"].resigned_on = date(2025, 1, 1)
    db.commit()
    roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})), today=date(2026, 10, 21))

    def emp(name):
        return db.scalar(select(Employee).where(Employee.name == name))

    assert str(emp("오일곱").resigned_on) == "2026-10-01"      # 인력현황 퇴사자 묶음의 퇴사일
    assert str(emp("떠난사람").resigned_on) == "2026-10-21"    # 인력현황에 없어 오늘 퇴사 처리
    assert emp("정다섯").resigned_on is None                    # 재직으로 되돌리면 지운다
    assert emp("이미퇴사").resigned_on is None                  # 예전 퇴사자는 건드리지 않는다


def test_화면에서_퇴사로_바꾸면_퇴사일이_남는다(admin_client, db):
    from datetime import date
    emp = Employee(emp_no="E1", name="가", status="ACTIVE")
    old = Employee(emp_no="E2", name="나", status="RESIGNED")
    db.add_all([emp, old])
    db.commit()
    admin_client.post(f"/employees/{emp.id}/edit", data={"emp_no": "E1", "name": "가", "status": "RESIGNED"})
    admin_client.post(f"/employees/{old.id}/edit", data={"emp_no": "E2", "name": "나", "status": "RESIGNED",
                                                         "note": "고침"})
    db.expire_all()
    assert db.get(Employee, emp.id).resigned_on == date.today()
    assert db.get(Employee, old.id).resigned_on is None          # 이미 퇴사였던 사람은 오늘로 찍지 않는다
    admin_client.post(f"/employees/{emp.id}/edit", data={"emp_no": "E1", "name": "가", "status": "ACTIVE"})
    db.expire_all()
    assert db.get(Employee, emp.id).resigned_on is None


# --- 인력현황 양식으로 내보내기 ---------------------------------------------------

def _export_rows(content: bytes) -> list[tuple]:
    from openpyxl import load_workbook
    return list(load_workbook(io.BytesIO(content), data_only=True).worksheets[0].iter_rows(values_only=True))


def test_내보내기는_원본과_같은_모양이다(db, staff):
    from datetime import date
    db.add(Asset(asset_no="A-2", name="모니터", status="IN_USE", holder_id=staff["김하나"].id))
    db.commit()
    roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})), today=date(2026, 10, 21))
    rows = _export_rows(roster.export_roster(db, today=date(2026, 10, 21)))

    assert rows[2][1] == "■ 인력현황"
    assert list(rows[3][1:]) == ["번호", "구분", "부서", "팀", "직급(책)", "성명", "고용형태", "입사일", "비고",
                                 "사번", "보유 자산", "자산번호"]
    names = [r[6] for r in rows if isinstance(r[1], int)]
    # 인력현황 순서 그대로, 협력사 묶음, 이번 달 퇴사자 순
    assert names[:5] == ["김하나", "이 둘", "박셋", "최넷", "정다섯"]
    titles = [r[1] for r in rows if isinstance(r[1], str) and r[1].startswith("*")]
    assert titles == ["* 협력사", "* 퇴사자"]
    kim = next(r for r in rows if r[6] == "김하나")
    assert (kim[2], kim[10], kim[11], kim[12]) == ("가법인", "E0001", 1, "A-2")
    leavers_header = next(r for r in rows if r[8] == "퇴사일")
    leavers = [r[6] for r in rows[rows.index(leavers_header) + 1:] if r[6]]
    assert sorted(leavers) == ["떠난사람", "오일곱"]           # 이번 달(10월) 퇴사자만


def test_지난달_퇴사자는_빠진다(db, staff):
    from datetime import date
    roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})), today=date(2026, 10, 21))
    rows = _export_rows(roster.export_roster(db, today=date(2026, 11, 2)))
    assert not any(r[1] == "* 퇴사자" for r in rows)
    rows = _export_rows(roster.export_roster(db, today=date(2026, 11, 2), month=date(2026, 10, 1)))
    assert any(r[1] == "* 퇴사자" for r in rows)


def test_내보낸_파일을_다시_올리면_바뀌는_것이_없다(db, staff):
    from datetime import date
    roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})), today=date(2026, 10, 21))
    content = roster.export_roster(db, today=date(2026, 10, 21))
    again = roster.read_roster(content)               # '요약' 시트는 건너뛰고 명단 시트를 고른다
    assert again.sheets == ["1021"]
    assert not roster.plan_sync(db, again).has_changes


def test_요약_시트(db, staff):
    from openpyxl import load_workbook
    roster.apply_sync(db, roster.read_roster(_file({"s": LATEST})))
    wb = load_workbook(io.BytesIO(roster.export_roster(db)), data_only=True)
    rows = list(wb["요약"].iter_rows(values_only=True))
    total = next(r for r in rows if r[0] == "합계")
    assert total[rows[0].index("합계")] == 6                 # 재직·휴직 6명 (협력사 포함)


def test_화면에서_내려받는다(admin_client, db, staff):
    response = admin_client.get("/employees/roster/export?month=2026-10")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "%EC%9D%B8%EB%A0%A5%ED%98%84%ED%99%A9" in response.headers["content-disposition"]   # 인력현황
    assert 'href="/employees/roster/export"' in admin_client.get("/employees").text


def test_조회_계정은_내보낼_수_없다(viewer_client):
    assert viewer_client.get("/employees/roster/export", follow_redirects=False).status_code == 403
