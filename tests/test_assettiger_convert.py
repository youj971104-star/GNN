"""AssetTiger 내보내기 파일 변환 테스트.

실제 회사 데이터 대신, 같은 형식의 견본 파일을 만들어 확인한다.
"""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from tools.convert_assettiger import convert, guess_category, looks_like_place, parse_money

TIGER_HEADERS = [
    "Asset Photo", "Asset Tag ID", "Description", "Purchase Date", "Cost",
    "Status", "Model", "Serial No", "Department", "Assigned to", "Event Notes",
]


def _tiger_file(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Asset"
    ws.append(TIGER_HEADERS)
    for row in rows:
        ws.append([row.get(h) for h in TIGER_HEADERS])
    wb.save(path)
    return path


def _row(tag, **cells):
    base = {
        "Asset Tag ID": tag,
        "Description": "노트북",
        "Purchase Date": datetime(2026, 2, 19),
        "Cost": "544,546.00",
        "Status": "Checked out",
        "Model": "V15 G4",
        "Serial No": f"SN-{tag}",
        "Department": "개발팀",
        "Assigned to": "홍길동",
        "Event Notes": "",
    }
    base.update(cells)
    return base


def _sheet_rows(path):
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    return headers, [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


# --- 값 변환 ------------------------------------------------------------------

@pytest.mark.parametrize(
    "description,expected",
    [
        ("노트북", "노트북"),
        ("데스크탑", "데스크톱"),
        ("27인치 모니터", "모니터"),
        ("파일 서버", "서버"),
        ("업무용 휴대폰", "모바일 기기"),
        ("Office 라이선스", "소프트웨어"),
        ("알 수 없는 물건", "기타"),
    ],
)
def test_설명에서_분류를_추정한다(description, expected):
    assert guess_category(description) == expected


def test_담당자_칸의_장소를_가려낸다():
    assert looks_like_place("서울사무실")
    assert looks_like_place("경기북부센터")
    assert looks_like_place("대전본사")
    assert not looks_like_place("홍길동")
    assert not looks_like_place("")


def test_쉼표가_있는_금액을_읽는다():
    assert parse_money("544,546.00") == 544546.0
    assert parse_money("1,502,719.00") == 1502719.0
    assert parse_money("") is None
    assert parse_money(None) is None


# --- 파일 변환 ----------------------------------------------------------------

def test_직원과_자산_파일이_만들어진다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [
        _row("0001", **{"Assigned to": "홍길동"}),
        _row("0002", **{"Assigned to": "김철수", "Department": "영업팀"}),
    ])
    info = convert(source, tmp_path / "out")

    assert info["rows"] == 2
    assert info["employees"] == 2
    assert info["employee_path"].exists()
    assert info["asset_path"].exists()


def test_담당자가_직원으로_만들어지고_자산에_연결된다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [_row("0001", **{"Assigned to": "홍길동"})])
    info = convert(source, tmp_path / "out")

    _, employees = _sheet_rows(info["employee_path"])
    assert employees[0]["이름"] == "홍길동"
    emp_no = employees[0]["사번"]

    _, assets = _sheet_rows(info["asset_path"])
    assert assets[0]["사용자 사번"] == emp_no   # 같은 사번으로 이어진다
    assert assets[0]["상태"] == "사용중"


def test_장소에_배정된_자산은_보관위치로_옮기고_재고로_둔다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [
        _row("0001", **{"Assigned to": "서울사무실", "Status": "Checked out"}),
    ])
    info = convert(source, tmp_path / "out")

    _, assets = _sheet_rows(info["asset_path"])
    assert assets[0]["보관위치"] == "서울사무실"
    assert assets[0]["상태"] == "재고"
    assert not assets[0]["사용자 사번"]

    # 장소는 직원으로 만들지 않는다
    _, employees = _sheet_rows(info["employee_path"])
    assert employees == []
    assert info["place_assigned"] == ["0001"]


def test_담당자가_없으면_재고로_둔다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [
        _row("0001", **{"Assigned to": "", "Status": "Available"}),
        # 담당자 없이 'Checked out' 인 모순된 행도 재고로 내린다
        _row("0002", **{"Assigned to": "", "Status": "Checked out"}),
    ])
    info = convert(source, tmp_path / "out")

    _, assets = _sheet_rows(info["asset_path"])
    assert all(row["상태"] == "재고" for row in assets)
    assert all(not row["사용자 사번"] for row in assets)


def test_모르는_상태값은_재고로_두고_알려준다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [
        _row("0001", **{"Status": "Something Odd", "Assigned to": ""}),
    ])
    info = convert(source, tmp_path / "out")

    assert info["unknown_status"] == ["Something Odd"]
    _, assets = _sheet_rows(info["asset_path"])
    assert assets[0]["상태"] == "재고"


def test_부서가_빈_행이_섞여도_채워진_값을_쓴다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [
        _row("0001", **{"Assigned to": "홍길동", "Department": ""}),
        _row("0002", **{"Assigned to": "홍길동", "Department": "개발팀"}),
    ])
    info = convert(source, tmp_path / "out")

    _, employees = _sheet_rows(info["employee_path"])
    assert len(employees) == 1
    assert employees[0]["부서"] == "개발팀"


def test_원본_메모를_비고에_남긴다(tmp_path):
    source = _tiger_file(tmp_path / "tiger.xlsx", [_row("0001", **{"Event Notes": "대표이사 노트북"})])
    info = convert(source, tmp_path / "out")

    _, assets = _sheet_rows(info["asset_path"])
    assert "대표이사 노트북" in assets[0]["비고"]
    assert "AssetTiger" in assets[0]["비고"]   # 어디서 왔는지도 남긴다


def test_변환한_파일을_그대로_시스템에_올릴_수_있다(tmp_path, db):
    """변환 결과가 실제 업로드 양식과 맞는지 확인한다."""
    from app.excel import import_assets, import_employees
    from app.models import Asset

    source = _tiger_file(tmp_path / "tiger.xlsx", [
        _row("0001", **{"Assigned to": "홍길동"}),
        _row("0002", **{"Assigned to": "", "Status": "Available"}),
    ])
    info = convert(source, tmp_path / "out")

    emp_result = import_employees(db, info["employee_path"].read_bytes())
    assert emp_result.skipped == 0 and emp_result.created == 1

    asset_result = import_assets(db, info["asset_path"].read_bytes(), actor="admin")
    assert asset_result.skipped == 0, asset_result.errors
    assert asset_result.created == 2

    from sqlalchemy import select
    asset = db.scalar(select(Asset).where(Asset.asset_no == "0001"))
    assert asset.status == "IN_USE"
    assert asset.holder.name == "홍길동"
    assert float(asset.purchase_price) == 544546.0
