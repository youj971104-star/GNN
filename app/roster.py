"""인력현황 엑셀을 직원 명부에 반영한다.

HR 이 관리하는 인력현황 파일(시트마다 한 시점)을 올리면 직원 명부를 그 시점에 맞춘다.

- 인력현황에 있고 명부에도 있는 사람: 법인만 채운다 (부서·직급 등은 건드리지 않는다)
- 인력현황에만 있는 사람: 새 직원으로 등록한다
- 명부에만 있는 사람(재직·휴직): 퇴사로 바꾼다
- 인력현황의 '* 퇴사자' 묶음에 있는 사람: 퇴사로 바꾼다

인력현황에는 사번이 없어 **이름으로** 맞춘다. 같은 이름이 둘 이상이면 어느 쪽인지 알 수
없으므로 건드리지 않고 '확인 필요' 로 따로 보여 준다.

먼저 미리보기(plan)로 바뀔 내용을 보여 주고, 확인을 받은 뒤 반영(apply)한다.
반영할 때는 그 시점의 명부로 다시 계산하므로, 그 사이에 명부가 바뀌어도 어긋나지 않는다.
"""

from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import numbering
from app.models import Asset, Employee

MAIN_GROUP = "본 인력현황"

# 머리글 칸 이름 (양식이 조금 달라도 읽을 수 있게 비슷한 이름을 함께 받는다)
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("성명", "이름"),
    "company": ("구분", "법인"),
    "department": ("부서",),
    "team": ("팀",),
    "position": ("직급(책)", "직급", "직책"),
    "kind": ("고용형태",),
    "date": ("입사일", "퇴사일"),
    "note": ("비고",),
}
HEADER_WORDS = {word for words in HEADER_ALIASES.values() for word in words} | {"번호"}

# 칸이 비어 있다는 표시로 쓰는 값
EMPTY_MARKS = {"", "-", "－", "—"}


class RosterError(ValueError):
    """인력현황 파일을 읽을 수 없을 때. 메시지를 그대로 화면에 보여 준다."""


@dataclass(frozen=True)
class RosterPerson:
    name: str
    company: str | None
    department: str | None
    team: str | None
    position: str | None
    kind: str | None              # 고용형태 (정규·계약·파견 …)
    joined: str | None            # 입사일(퇴사자 묶음은 퇴사일), YYYY-MM-DD
    note: str | None
    group: str                    # 묶음 이름 (본 인력현황 / 협력사 / 퇴사자 …)
    resigned: bool
    row: int

    @property
    def key(self) -> str:
        return name_key(self.name)

    @property
    def on_leave(self) -> bool:
        return bool(self.note and "휴직" in self.note)


@dataclass
class Roster:
    sheets: list[str]
    sheet: str
    people: list[RosterPerson]

    @property
    def active(self) -> list[RosterPerson]:
        return [p for p in self.people if not p.resigned]

    @property
    def resigned(self) -> list[RosterPerson]:
        return [p for p in self.people if p.resigned]

    def group_counts(self) -> list[tuple[str, int]]:
        counts = Counter(p.group for p in self.people)
        return list(counts.items())

    def company_counts(self) -> list[tuple[str, int]]:
        counts = Counter(p.company or "(비어 있음)" for p in self.active)
        return counts.most_common()

    # 미리보기 → 반영 사이에 화면에 실어 보내기 위한 변환
    def to_data(self) -> dict:
        return {"sheets": self.sheets, "sheet": self.sheet, "people": [asdict(p) for p in self.people]}

    @classmethod
    def from_data(cls, data: dict) -> "Roster":
        return cls(
            sheets=list(data["sheets"]),
            sheet=data["sheet"],
            people=[RosterPerson(**p) for p in data["people"]],
        )


def name_key(name: str) -> str:
    """이름 비교용: 공백과 괄호 속 설명을 없앤다.

    '홍 길동' = '홍길동' = '홍길동(QA1팀)'. 예전 시스템에서 옮겨 온 이름에 부서 등이
    괄호로 붙어 있으면, 같은 사람이 '퇴사 처리 + 신규 등록' 으로 갈라지지 않게 한다.
    """
    without_notes = re.sub(r"[(（\[].*?[)）\]]", "", name or "")
    return re.sub(r"\s+", "", without_notes)


def _text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return None if text in EMPTY_MARKS else text


def _header_map(cells: list) -> dict[str, int] | None:
    """머리글 줄이면 {항목: 칸 위치}, 아니면 None. '성명'과 '구분' 이 함께 있어야 머리글로 본다."""
    texts = [(_text(c) or "") for c in cells]
    mapping: dict[str, int] = {}
    for field_name, aliases in HEADER_ALIASES.items():
        for pos, text in enumerate(texts):
            if text in aliases:
                mapping[field_name] = pos
                break
    if "name" in mapping and "company" in mapping:
        return mapping
    return None


def _group_title(cells: list) -> str | None:
    """'* 협력사' · '* 퇴사자' 같은 묶음 제목이면 그 이름."""
    for cell in cells:
        text = _text(cell)
        if text and text.startswith("*"):
            return text.lstrip("*").strip() or None
    return None


def read_roster(content: bytes, sheet: str | None = None) -> Roster:
    """인력현황 엑셀을 읽는다. 시트를 지정하지 않으면 맨 오른쪽(가장 최근) 시트."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # 엑셀이 아닌 파일
        raise RosterError(f"엑셀 파일을 열 수 없습니다: {exc}") from exc

    try:
        sheets = list(wb.sheetnames)
        if not sheets:
            raise RosterError("시트가 없는 파일입니다.")
        name = sheet if sheet in sheets else sheets[-1]
        ws = wb[name]

        people: list[RosterPerson] = []
        mapping: dict[str, int] | None = None
        group = MAIN_GROUP
        for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = list(row)
            title = _group_title(cells)
            if title:
                group = title
                mapping = None          # 묶음마다 머리글이 다시 나온다
                continue
            header = _header_map(cells)
            if header:
                mapping = header
                continue
            if mapping is None:
                continue

            def cell(field_name: str):
                pos = mapping.get(field_name)
                return _text(cells[pos]) if pos is not None and pos < len(cells) else None

            person_name = cell("name")
            if not person_name or person_name in HEADER_WORDS:
                continue
            people.append(
                RosterPerson(
                    name=person_name,
                    company=cell("company"),
                    department=cell("department"),
                    team=cell("team"),
                    position=cell("position"),
                    kind=cell("kind"),
                    joined=cell("date"),
                    note=cell("note"),
                    group=group,
                    resigned="퇴사" in group,
                    row=row_no,
                )
            )
    finally:
        wb.close()

    if not people:
        raise RosterError(
            f"'{name}' 시트에서 직원을 찾지 못했습니다. "
            "'성명'과 '구분' 칸이 있는 인력현황 양식인지 확인해 주세요."
        )
    return Roster(sheets=sheets, sheet=name, people=people)


# --- 미리보기 --------------------------------------------------------------------

@dataclass
class Match:
    employee: Employee
    person: RosterPerson
    old_company: str | None = None


@dataclass
class Resign:
    employee: Employee
    reason: str
    holding: int = 0
    person: RosterPerson | None = None


@dataclass
class Ambiguous:
    name: str
    employees: list[Employee]
    people: list[RosterPerson]


@dataclass
class SyncPlan:
    company_changes: list[Match] = field(default_factory=list)   # 법인을 새로 넣거나 바꿈
    unchanged: list[Match] = field(default_factory=list)          # 법인이 이미 같음
    reactivate: list[Match] = field(default_factory=list)         # 명부에선 퇴사인데 인력현황엔 재직
    add: list[RosterPerson] = field(default_factory=list)
    resign: list[Resign] = field(default_factory=list)
    ambiguous: list[Ambiguous] = field(default_factory=list)
    unknown_resigned: list[RosterPerson] = field(default_factory=list)  # 퇴사자 명단에 있지만 명부에 없음

    @property
    def resign_holding(self) -> list[Resign]:
        return [r for r in self.resign if r.holding]

    @property
    def has_changes(self) -> bool:
        return bool(self.company_changes or self.reactivate or self.add or self.resign)


def plan_sync(db: Session, roster: Roster) -> SyncPlan:
    plan = SyncPlan()

    employees = list(db.scalars(select(Employee).order_by(Employee.name, Employee.id)).all())
    by_key: dict[str, list[Employee]] = defaultdict(list)
    for emp in employees:
        by_key[name_key(emp.name)].append(emp)

    holding = dict(
        db.execute(
            select(Asset.holder_id, func.count(Asset.id))
            .where(Asset.holder_id.is_not(None))
            .group_by(Asset.holder_id)
        ).all()
    )

    active: dict[str, list[RosterPerson]] = defaultdict(list)
    for person in roster.active:
        active[person.key].append(person)
    resigned: dict[str, list[RosterPerson]] = defaultdict(list)
    for person in roster.resigned:
        if person.key not in active:           # 다시 입사해 재직 명단에도 있으면 재직이 우선
            resigned[person.key].append(person)

    handled: set[int] = set()                  # 인력현황에서 찾은(또는 판단을 미룬) 직원

    for key, people in active.items():
        matches = by_key.get(key, [])
        if len(people) > 1 or len(matches) > 1:
            plan.ambiguous.append(Ambiguous(people[0].name, matches, people))
            handled.update(emp.id for emp in matches)
            continue
        person = people[0]
        if not matches:
            plan.add.append(person)
            continue
        emp = matches[0]
        handled.add(emp.id)
        match = Match(emp, person, emp.company)
        if emp.status == "RESIGNED":
            plan.reactivate.append(match)
        elif (emp.company or None) != person.company:
            plan.company_changes.append(match)
        else:
            plan.unchanged.append(match)

    for key, people in resigned.items():
        matches = by_key.get(key, [])
        if len(people) > 1 or len(matches) > 1:
            plan.ambiguous.append(Ambiguous(people[0].name, matches, people))
            handled.update(emp.id for emp in matches)
            continue
        if not matches:
            plan.unknown_resigned.append(people[0])
            continue
        emp = matches[0]
        handled.add(emp.id)
        if emp.status != "RESIGNED":
            plan.resign.append(Resign(emp, "인력현황 퇴사자 명단", holding.get(emp.id, 0), people[0]))
        elif (emp.company or None) != people[0].company:
            plan.company_changes.append(Match(emp, people[0], emp.company))

    for emp in employees:
        if emp.id in handled or emp.status == "RESIGNED":
            continue
        plan.resign.append(Resign(emp, "인력현황에 없음", holding.get(emp.id, 0)))

    plan.resign.sort(key=lambda r: (-r.holding, r.employee.name))
    return plan


# --- 반영 ------------------------------------------------------------------------

@dataclass
class SyncResult:
    company: int = 0
    reactivated: int = 0
    added: int = 0
    resigned: int = 0
    resigned_holding: int = 0
    ambiguous: int = 0

    def summary(self) -> str:
        parts = [
            f"법인 입력 {self.company}명",
            f"신규 등록 {self.added}명",
            f"퇴사 처리 {self.resigned}명",
        ]
        if self.reactivated:
            parts.append(f"재직으로 되돌림 {self.reactivated}명")
        text = "인력현황을 반영했습니다: " + ", ".join(parts) + "."
        if self.resigned_holding:
            text += f" 퇴사 처리한 {self.resigned_holding}명은 아직 자산을 갖고 있으니 회수해 주세요."
        if self.ambiguous:
            text += f" 같은 이름이 겹친 {self.ambiguous}건은 건드리지 않았습니다."
        return text


def _new_employee_note(person: RosterPerson, sheet: str) -> str:
    parts = [f"인력현황({sheet})에서 등록"]
    if person.group != MAIN_GROUP:
        parts.append(person.group)
    if person.kind:
        parts.append(f"고용형태 {person.kind}")
    if person.team and person.team != person.department:
        parts.append(f"팀 {person.team}")
    if person.joined:
        parts.append(f"입사일 {person.joined}")
    if person.note:
        parts.append(person.note)
    return " · ".join(parts)


def apply_sync(db: Session, roster: Roster) -> SyncResult:
    """지금 명부 기준으로 다시 계산해 한 번에 반영한다 (중간에 실패하면 아무것도 바뀌지 않는다)."""
    plan = plan_sync(db, roster)
    result = SyncResult(ambiguous=len(plan.ambiguous))
    try:
        for match in plan.company_changes:
            match.employee.company = match.person.company
            result.company += 1
        for match in plan.reactivate:
            match.employee.company = match.person.company
            match.employee.status = "LEAVE" if match.person.on_leave else "ACTIVE"
            result.reactivated += 1
        numbers = numbering.issue_employee_numbers(db, len(plan.add))
        for person, emp_no in zip(plan.add, numbers):
            db.add(
                Employee(
                    emp_no=emp_no,
                    name=person.name,
                    company=person.company,
                    department=person.department,
                    position=person.position,
                    status="LEAVE" if person.on_leave else "ACTIVE",
                    note=_new_employee_note(person, roster.sheet),
                )
            )
            result.added += 1
        for item in plan.resign:
            item.employee.status = "RESIGNED"
            if item.person is not None and item.person.company:
                item.employee.company = item.person.company
            result.resigned += 1
            result.resigned_holding += int(bool(item.holding))
        db.commit()
    except Exception:
        db.rollback()
        raise
    return result
